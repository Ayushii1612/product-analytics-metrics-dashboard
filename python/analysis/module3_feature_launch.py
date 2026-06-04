"""
python/analysis/module3_feature_launch.py

Module 3: Feature Launch Impact Analysis
  • Pre vs post-launch comparison (new checkout flow, 2024-05-01)
  • Primary metric: Conversion Rate
  • Guardrail metrics: page_load_ms, bounce_rate, revenue_per_session
  • Statistical significance testing with confidence intervals
  • Side effect detection on secondary metrics
  • Launch recommendation report
"""

import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

import pandas as pd
import numpy as np
from pathlib import Path
from scipy import stats
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import SIMULATION


# ─── Load Data ───────────────────────────────────────────
def load_data() -> pd.DataFrame:
    path = Path(__file__).parent.parent.parent / "data" / "raw" / "sessions.parquet"
    df = pd.read_parquet(path)
    df["timestamp"] = pd.to_datetime(df["timestamp"])
    df["date"]      = pd.to_datetime(df["date"])
    return df


# ─── Helpers ─────────────────────────────────────────────
def two_proportion_ztest(n1, conv1, n2, conv2):
    """Returns z_stat, p_value, ci_low, ci_high for difference in proportions."""
    from statsmodels.stats.proportion import proportions_ztest, proportion_confint
    p1, p2 = conv1/n1, conv2/n2
    z_stat, p_value = proportions_ztest([conv1, conv2], [n1, n2])
    # 95% CI for difference using normal approximation
    se = np.sqrt(p1*(1-p1)/n1 + p2*(1-p2)/n2)
    diff = p2 - p1
    ci_low  = diff - 1.96 * se
    ci_high = diff + 1.96 * se
    return z_stat, p_value, ci_low, ci_high


def welch_ttest(arr1: pd.Series, arr2: pd.Series):
    """Welch t-test for continuous metrics."""
    t_stat, p_value = stats.ttest_ind(arr1, arr2, equal_var=False)
    n1, n2   = len(arr1), len(arr2)
    m1, m2   = arr1.mean(), arr2.mean()
    se       = np.sqrt(arr1.var(ddof=1)/n1 + arr2.var(ddof=1)/n2)
    ci_low   = (m2 - m1) - 1.96 * se
    ci_high  = (m2 - m1) + 1.96 * se
    return t_stat, p_value, m1, m2, ci_low, ci_high


# ─── Pre / Post Datasets ─────────────────────────────────
def split_pre_post(df: pd.DataFrame, launch_date: str, window_days: int = 30):
    """
    Returns (df_pre, df_post) using an equal window before and after launch.
    Excludes the drop event window to avoid confounding.
    """
    launch_dt = pd.Timestamp(launch_date)
    drop_dt   = pd.Timestamp(SIMULATION["drop_event_date"])
    drop_end  = drop_dt + pd.Timedelta(days=SIMULATION["drop_duration_days"] + 1)

    pre_start = launch_dt - pd.Timedelta(days=window_days)
    post_end  = launch_dt + pd.Timedelta(days=window_days)

    df_pre  = df[(df["date"] >= pre_start) & (df["date"] < launch_dt) &
                 ~((df["date"] >= drop_dt) & (df["date"] < drop_end))]
    df_post = df[(df["date"] >= launch_dt) & (df["date"] < post_end)]

    return df_pre, df_post


# ─── Primary Metric: Conversion Rate ─────────────────────
def analyze_conversion_rate(df_pre, df_post) -> dict:
    n1    = len(df_pre)
    n2    = len(df_post)
    conv1 = df_pre["converted"].sum()
    conv2 = df_post["converted"].sum()
    cr1   = conv1 / n1
    cr2   = conv2 / n2

    z_stat, p_value, ci_low, ci_high = two_proportion_ztest(n1, conv1, n2, conv2)
    relative_lift = (cr2 - cr1) / cr1

    return {
        "metric"             : "Conversion Rate",
        "type"               : "primary",
        "pre_value"          : round(cr1 * 100, 4),
        "post_value"         : round(cr2 * 100, 4),
        "absolute_change_pp" : round((cr2 - cr1) * 100, 4),
        "relative_lift_pct"  : round(relative_lift * 100, 2),
        "ci_95_low_pp"       : round(ci_low * 100, 4),
        "ci_95_high_pp"      : round(ci_high * 100, 4),
        "z_statistic"        : round(z_stat, 4),
        "p_value"            : round(p_value, 8),
        "is_significant"     : p_value < 0.05,
        "pre_sessions"       : n1,
        "post_sessions"      : n2,
        "verdict"            : "SIGNIFICANT IMPROVEMENT ✅" if (p_value < 0.05 and cr2 > cr1)
                               else ("SIGNIFICANT REGRESSION ⚠️" if p_value < 0.05
                               else "NOT SIGNIFICANT — needs more data 🔄"),
    }


# ─── Guardrail Metrics ────────────────────────────────────
def analyze_guardrails(df_pre, df_post) -> list[dict]:
    guardrails = []

    # 1. Page Load Time
    t_stat, p_val, m1, m2, ci_low, ci_high = welch_ttest(
        df_pre["page_load_ms"], df_post["page_load_ms"]
    )
    guardrails.append({
        "metric"             : "Page Load Time (ms)",
        "type"               : "guardrail",
        "direction_good"     : "lower",
        "pre_value"          : round(m1, 1),
        "post_value"         : round(m2, 1),
        "absolute_change"    : round(m2 - m1, 1),
        "relative_change_pct": round((m2 - m1) / m1 * 100, 2),
        "ci_95_low"          : round(ci_low, 1),
        "ci_95_high"         : round(ci_high, 1),
        "p_value"            : round(p_val, 6),
        "is_significant"     : p_val < 0.05,
        "passed"             : m2 <= m1,
        "verdict"            : "✅ PASSED — Load time improved" if m2 < m1
                               else "⚠️ FAILED — Load time degraded",
    })

    # 2. Bounce Rate
    br1 = df_pre["bounced"].mean()
    br2 = df_post["bounced"].mean()
    n1, n2 = len(df_pre), len(df_post)
    z_stat, p_val, ci_low, ci_high = two_proportion_ztest(
        n1, df_pre["bounced"].sum(), n2, df_post["bounced"].sum()
    )
    guardrails.append({
        "metric"             : "Bounce Rate",
        "type"               : "guardrail",
        "direction_good"     : "lower",
        "pre_value"          : round(br1 * 100, 3),
        "post_value"         : round(br2 * 100, 3),
        "absolute_change_pp" : round((br2 - br1) * 100, 3),
        "relative_change_pct": round((br2 - br1) / br1 * 100, 2),
        "p_value"            : round(p_val, 6),
        "is_significant"     : p_val < 0.05,
        "passed"             : br2 <= br1,
        "verdict"            : "✅ PASSED — Bounce rate held/improved" if br2 <= br1
                               else "⚠️ FAILED — Bounce rate increased",
    })

    # 3. Revenue Per Session
    t_stat, p_val, m1, m2, ci_low, ci_high = welch_ttest(
        df_pre["revenue"], df_post["revenue"]
    )
    guardrails.append({
        "metric"             : "Revenue Per Session (₹)",
        "type"               : "guardrail",
        "direction_good"     : "higher",
        "pre_value"          : round(m1, 2),
        "post_value"         : round(m2, 2),
        "absolute_change"    : round(m2 - m1, 2),
        "relative_change_pct": round((m2 - m1) / m1 * 100, 2),
        "ci_95_low"          : round(ci_low, 2),
        "ci_95_high"         : round(ci_high, 2),
        "p_value"            : round(p_val, 6),
        "is_significant"     : p_val < 0.05,
        "passed"             : m2 >= m1,
        "verdict"            : "✅ PASSED — Revenue per session improved" if m2 >= m1
                               else "⚠️ FAILED — Revenue per session dropped",
    })

    return guardrails


# ─── Side Effects: Secondary Metrics ─────────────────────
def analyze_side_effects(df_pre, df_post) -> list[dict]:
    effects = []

    # Check conversion by device
    for device in ["mobile", "desktop", "tablet"]:
        pre_d  = df_pre[df_pre["device"] == device]
        post_d = df_post[df_post["device"] == device]
        if len(pre_d) < 100 or len(post_d) < 100:
            continue
        cr_pre  = pre_d["converted"].mean()
        cr_post = post_d["converted"].mean()
        lift    = (cr_post - cr_pre) / cr_pre
        effects.append({
            "segment"   : f"Device: {device}",
            "pre_cr"    : round(cr_pre * 100, 4),
            "post_cr"   : round(cr_post * 100, 4),
            "lift_pct"  : round(lift * 100, 2),
            "concern"   : "⚠️ Regression on segment" if lift < -0.05 else "✅ OK",
        })

    # Check by traffic source
    for source in df_pre["traffic_source"].unique():
        pre_s  = df_pre[df_pre["traffic_source"] == source]
        post_s = df_post[df_post["traffic_source"] == source]
        if len(pre_s) < 500 or len(post_s) < 500:
            continue
        cr_pre  = pre_s["converted"].mean()
        cr_post = post_s["converted"].mean()
        lift    = (cr_post - cr_pre) / cr_pre
        effects.append({
            "segment"  : f"Source: {source}",
            "pre_cr"   : round(cr_pre * 100, 4),
            "post_cr"  : round(cr_post * 100, 4),
            "lift_pct" : round(lift * 100, 2),
            "concern"  : "⚠️ Regression on segment" if lift < -0.05 else "✅ OK",
        })

    return effects


# ─── Launch Recommendation Report ─────────────────────────
def generate_launch_report(cr_result, guardrails, side_effects) -> str:
    guardrail_pass_count = sum(1 for g in guardrails if g["passed"])
    all_guardrails_pass  = guardrail_pass_count == len(guardrails)
    primary_significant  = cr_result["is_significant"] and cr_result["post_value"] > cr_result["pre_value"]

    if primary_significant and all_guardrails_pass:
        recommendation = "🚀 SHIP IT — Full rollout recommended"
        rationale = (
            "The new checkout feature shows a statistically significant conversion "
            "rate improvement with no guardrail metric regressions."
        )
    elif primary_significant and not all_guardrails_pass:
        recommendation = "⚠️ CONDITIONAL SHIP — Fix guardrail issues before full rollout"
        rationale = (
            "Primary metric improved significantly, but one or more guardrail metrics "
            "show concerning signals. Investigate and fix before scaling."
        )
    else:
        recommendation = "🛑 DO NOT SHIP — No significant improvement detected"
        rationale = "Insufficient evidence of improvement. Continue testing or redesign."

    lines = [
        "=" * 70,
        "   FEATURE LAUNCH IMPACT ANALYSIS REPORT",
        f"   Feature: New Checkout Flow | Launch: {SIMULATION['feature_launch_date']}",
        "=" * 70,
        "",
        f"RECOMMENDATION: {recommendation}",
        f"RATIONALE: {rationale}",
        "",
        "─" * 60,
        "PRIMARY METRIC: CONVERSION RATE",
        "─" * 60,
        f"  Pre-launch CR        : {cr_result['pre_value']}%",
        f"  Post-launch CR       : {cr_result['post_value']}%",
        f"  Absolute Change      : {cr_result['absolute_change_pp']:+.4f}pp",
        f"  Relative Lift        : {cr_result['relative_lift_pct']:+.2f}%",
        f"  95% CI               : [{cr_result['ci_95_low_pp']:+.4f}pp, {cr_result['ci_95_high_pp']:+.4f}pp]",
        f"  Z-Statistic          : {cr_result['z_statistic']}",
        f"  P-Value              : {cr_result['p_value']}",
        f"  VERDICT              : {cr_result['verdict']}",
        "",
        "─" * 60,
        "GUARDRAIL METRICS",
        "─" * 60,
    ]
    for g in guardrails:
        lines.append(f"  {g['metric']:<35} {g['verdict']}")
        if "absolute_change_pp" in g:
            lines.append(f"    Pre: {g['pre_value']} → Post: {g['post_value']}  ({g['absolute_change_pp']:+.3f}pp)")
        else:
            lines.append(f"    Pre: {g['pre_value']} → Post: {g['post_value']}  ({g['relative_change_pct']:+.2f}%)")

    lines += [
        "",
        "─" * 60,
        "SIDE EFFECTS — SEGMENT-LEVEL ANALYSIS",
        "─" * 60,
    ]
    for e in side_effects:
        lines.append(f"  {e['segment']:<35} Lift: {e['lift_pct']:+.2f}%  {e['concern']}")

    lines += [
        "",
        "─" * 60,
        "NEXT STEPS",
        "─" * 60,
        "  1. Monitor conversion rate for 2 more weeks post-launch",
        "  2. Set up automated alerts for >2σ deviation in daily CR",
        "  3. Investigate mobile segment for any remaining gap vs desktop",
        "  4. Run qualitative research (session recordings) on new checkout flow",
        "  5. Track revenue per session weekly — watch for AOV changes",
        "=" * 70,
    ]
    return "\n".join(lines)


def save_report(report_text, cr_result, guardrails, side_effects):
    out_dir = Path(__file__).parent.parent.parent / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)

    txt_path = out_dir / "module3_feature_launch_report.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(report_text)
    print(f"  ✓ Text report → {txt_path}")

    # Excel
    xlsx_path = out_dir / "module3_feature_launch_report.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Launch Impact"

    dark_fill  = PatternFill("solid", fgColor="1F3864")
    green_fill = PatternFill("solid", fgColor="E2EFDA")
    red_fill   = PatternFill("solid", fgColor="FCE4D6")

    ws.merge_cells("A1:F1")
    ws["A1"] = "FEATURE LAUNCH IMPACT ANALYSIS — New Checkout Flow"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = dark_fill
    ws.row_dimensions[1].height = 26

    # Primary metric
    ws["A3"] = "PRIMARY METRIC"
    ws["A3"].font = Font(bold=True, size=11)
    headers = ["Metric", "Pre-Launch", "Post-Launch", "Absolute Δ", "Relative Δ", "Verdict"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=4, column=col, value=h).font = Font(bold=True)
    ws.cell(row=5, column=1, value=cr_result["metric"])
    ws.cell(row=5, column=2, value=f"{cr_result['pre_value']}%")
    ws.cell(row=5, column=3, value=f"{cr_result['post_value']}%")
    ws.cell(row=5, column=4, value=f"{cr_result['absolute_change_pp']:+.4f}pp")
    ws.cell(row=5, column=5, value=f"{cr_result['relative_lift_pct']:+.2f}%")
    ws.cell(row=5, column=6, value=cr_result["verdict"])
    fill = green_fill if cr_result["is_significant"] else red_fill
    for col in range(1, 7):
        ws.cell(row=5, column=col).fill = fill

    # Guardrails
    ws["A7"] = "GUARDRAIL METRICS"
    ws["A7"].font = Font(bold=True, size=11)
    for col, h in enumerate(headers, 1):
        ws.cell(row=8, column=col, value=h).font = Font(bold=True)
    for i, g in enumerate(guardrails, start=9):
        ws.cell(row=i, column=1, value=g["metric"])
        ws.cell(row=i, column=2, value=str(g["pre_value"]))
        ws.cell(row=i, column=3, value=str(g["post_value"]))
        ws.cell(row=i, column=6, value=g["verdict"])
        fill = green_fill if g["passed"] else red_fill
        for col in range(1, 7):
            ws.cell(row=i, column=col).fill = fill

    for col, width in zip(range(1, 7), [35, 15, 15, 15, 15, 40]):
        ws.column_dimensions[get_column_letter(col)].width = width

    # Sheet 2: Segment side effects
    ws2 = wb.create_sheet("Segment Analysis")
    ws2["A1"] = "Segment-Level Side Effects"
    ws2["A1"].font = Font(bold=True, size=12)
    for col, h in enumerate(["Segment", "Pre CR%", "Post CR%", "Lift%", "Status"], 1):
        ws2.cell(row=3, column=col, value=h).font = Font(bold=True)
    for i, e in enumerate(side_effects, start=4):
        ws2.cell(row=i, column=1, value=e["segment"])
        ws2.cell(row=i, column=2, value=e["pre_cr"])
        ws2.cell(row=i, column=3, value=e["post_cr"])
        ws2.cell(row=i, column=4, value=e["lift_pct"])
        ws2.cell(row=i, column=5, value=e["concern"])
        fill = green_fill if "✅" in e["concern"] else red_fill
        for col in range(1, 6):
            ws2.cell(row=i, column=col).fill = fill
    for col, width in zip(range(1, 6), [40, 12, 12, 12, 35]):
        ws2.column_dimensions[get_column_letter(col)].width = width

    wb.save(xlsx_path)
    print(f"  ✓ Excel report → {xlsx_path}")


# ─── Main ────────────────────────────────────────────────
def run():
    print("\n🚀  Module 3: Feature Launch Impact Analysis")
    print("─" * 50)

    df = load_data()
    launch_date = SIMULATION["feature_launch_date"]
    print(f"  Feature launch date: {launch_date}")

    df_pre, df_post = split_pre_post(df, launch_date, window_days=30)
    print(f"  Pre-launch sessions : {len(df_pre):,}")
    print(f"  Post-launch sessions: {len(df_post):,}")

    cr_result   = analyze_conversion_rate(df_pre, df_post)
    guardrails  = analyze_guardrails(df_pre, df_post)
    side_effects = analyze_side_effects(df_pre, df_post)

    print(f"\n  Primary metric: CR {cr_result['pre_value']}% → {cr_result['post_value']}%")
    print(f"  Lift: {cr_result['relative_lift_pct']:+.2f}%  |  {cr_result['verdict']}")
    print("\n  Guardrail results:")
    for g in guardrails:
        print(f"    {g['metric']:<35} {g['verdict']}")

    report = generate_launch_report(cr_result, guardrails, side_effects)
    save_report(report, cr_result, guardrails, side_effects)

    print("\n✅  Module 3 complete!\n")
    return cr_result, guardrails, side_effects


if __name__ == "__main__":
    run()